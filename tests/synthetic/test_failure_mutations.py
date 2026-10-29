from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image


TEST_ROOT = Path(__file__).resolve().parent
REPO_ROOT = TEST_ROOT.parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))
sys.path.insert(0, str(TEST_ROOT))

from application_pack.artifacts import canonical_json, sha256_bytes, sha256_file  # noqa: E402
from application_pack.contracts import ContractError  # noqa: E402
from application_pack.engine import run_manifest, validate_run  # noqa: E402
from application_pack.invariant_checks import atomic_phrase_findings  # noqa: E402
from application_pack.rendering import render_exact_artifact  # noqa: E402
from test_engine import make_fixture, write_json  # noqa: E402


CATALOG = json.loads((TEST_ROOT / "failure_catalog.json").read_text(encoding="utf-8"))


class FailureMutationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.manifest = make_fixture(self.root)

    def tearDown(self) -> None:
        self.temp.cleanup()

    @property
    def run_dir(self) -> Path:
        return self.root / "runs" / "synthetic-run"

    def load(self, path: Path) -> dict:
        return json.loads(path.read_text(encoding="utf-8"))

    def save(self, path: Path, value: dict) -> None:
        write_json(path, value)

    def gate_status(self, report: dict, gate_id: str) -> str | None:
        return next((gate["status"] for gate in report["gates"] if gate["id"] == gate_id), None)

    def run_failure_mutation(self, entry: dict) -> None:
        mutation = entry["mutation"]
        stages = self.root / "stages"

        manual_checks = {
            "fail_manual_full_page_readability": ("full_page_readability", "layout_quality"),
            "fail_manual_physical_edge_clearance": ("physical_edge_clearance", "layout_quality"),
            "fail_manual_natural_content_flow": ("natural_content_flow", "layout_quality"),
            "fail_manual_metric_visibility": ("metric_visibility", "layout_quality"),
            "fail_manual_continuous_page_review": ("continuous_page_review", "layout_quality"),
            "fail_manual_talk_taxonomy": ("talk_taxonomy", "career_provenance"),
            "fail_manual_current_link_inventory": ("current_link_inventory", "cross_format_parity"),
            "fail_manual_docx_hanging_indents": ("docx_hanging_indents", "layout_quality"),
            "fail_manual_role_date_rows_atomic": ("role_date_rows_atomic", "layout_quality"),
            "fail_manual_source_emphasis_preservation": ("source_emphasis_preservation", "layout_quality"),
        }
        if mutation in manual_checks:
            check_name, gate_id = manual_checks[mutation]
            path = stages / "review-record.json"
            data = self.load(path)
            data["manual_checks"][check_name] = False
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, gate_id), "fail", entry)
        elif mutation == "remove_baseline_source":
            path = stages / "evidence-bundle.json"
            data = self.load(path)
            data["sources"] = [source for source in data["sources"] if source["kind"] != "baseline"]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "evidence_integrity"), "fail", entry)
        elif mutation == "omit_coverage_entry":
            path = stages / "coverage-ledger.json"
            data = self.load(path)
            data["entries"] = [item for item in data["entries"] if item["claim_id"] != "research-claim"]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "substance_preservation"), "fail", entry)
        elif mutation == "weaken_protected_claim":
            path = stages / "coverage-ledger.json"
            data = self.load(path)
            next(item for item in data["entries"] if item["claim_id"] == "platform-claim")["strength_preserved"] = False
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "substance_preservation"), "fail", entry)
        elif mutation == "wrong_page_count":
            data = self.load(self.manifest)
            data["constraints"]["page_counts"]["resume"] = 3
            self.save(self.manifest, data)
            with self.assertRaises(ContractError, msg=entry):
                run_manifest(self.manifest)
        elif mutation == "duplicate_outline_section":
            path = stages / "resume.json"
            data = self.load(path)
            data["sections"].append({"id": data["sections"][0]["id"], "heading": "Later experience"})
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "persuasive_content"), "fail", entry)
        elif mutation == "duplicate_claim":
            path = stages / "resume.json"
            data = self.load(path)
            data["claims"].append(dict(data["claims"][0]))
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "persuasive_content"), "fail", entry)
        elif mutation == "repeat_opening_verb":
            path = stages / "resume.json"
            data = self.load(path)
            for claim in data["claims"]:
                claim["text"] = f"Led {claim['text'].lower()}"
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "persuasive_content"), "fail", entry)
        elif mutation == "purposeless_element":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["visible_elements"][0]["reader_purpose"] = ""
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "duplicate_identity_block":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["visible_elements"].append({
                "document_id": "resume", "element_id": "identity-again", "function": "identity",
                "reader_purpose": "Repeat the candidate name",
            })
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "misclassify_salutation":
            path = stages / "design-contract.json"
            data = self.load(path)
            next(item for item in data["visible_elements"] if item["element_id"] == "salutation")["function"] = "body"
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "drop_review_invalidation":
            path = stages / "change-set.json"
            data = self.load(path)
            data["changes"] = [{"id": "change-one", "target_ids": ["identity"], "invalidates": ["render", "release"]}]
            data["visual_evidence"]["reviewed_change_ids"] = ["change-one"]
            data["visual_evidence"]["after_page_hashes"] = [
                self.load(stages / "render-evidence.json")["artifacts"][0]["pages"][0]["sha256"]
            ]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "change_isolation"), "fail", entry)
        elif mutation == "unrequested_tagline":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["page_constraints"]["resume_tagline"] = False
            data["visible_elements"].append({
                "document_id": "resume", "element_id": "tagline", "function": "tagline",
                "reader_purpose": "Repeat positioning",
            })
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "unjustified_page_break":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["page_breaks"] = [{
                "document_id": "resume", "after_element_id": "identity", "reason": " ",
                "previous_fill_ratio": 0.9, "next_block_fraction": 0.2,
            }]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "page_number":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["visible_elements"].append({
                "document_id": "resume",
                "element_id": "page-number",
                "function": "page_number",
                "reader_purpose": "Show the page",
            })
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "bad_provenance":
            path = stages / "cover-letter.json"
            data = self.load(path)
            data["claims"][0]["parent_entity_id"] = "independent"
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "career_provenance"), "fail", entry)
        elif mutation == "stale_positioning_inputs":
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 0, report)
            path = self.run_dir / "work" / "state.json"
            data = self.load(path)
            data["stages"]["positioning"]["input_hashes"] = {}
            self.save(path, data)
            report = validate_run(self.run_dir)
            self.assertEqual(self.gate_status(report, "state_freshness"), "fail", entry)
        elif mutation == "accepted_region_hash_mismatch":
            path = stages / "change-set.json"
            data = self.load(path)
            data["accepted_regions"] = [{
                "artifact_id": "resume_page_1", "sha256": "0" * 64, "description": "Accepted first page"
            }]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "change_isolation"), "fail", entry)
        elif mutation == "stale_review_hash":
            path = stages / "review-record.json"
            data = self.load(path)
            data["artifact_hashes"]["resume_pdf"] = "0" * 64
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "exact_artifact_review"), "fail", entry)
        elif mutation == "collateral_change":
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 0, report)
            path = self.run_dir / "audit" / "change-set.json"
            data = self.load(path)
            data["collateral_change_ids"] = ["accepted-page-two"]
            self.save(path, data)
            report = validate_run(self.run_dir)
            self.assertEqual(self.gate_status(report, "change_isolation"), "fail", entry)
        elif mutation == "failed_gate_no_promotion":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["visible_elements"][0]["reader_purpose"] = ""
            self.save(path, data)
            code, _ = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertFalse((self.run_dir / "final").exists(), entry)
        elif mutation == "duplicate_semantic_style":
            path = stages / "design-contract.json"
            data = self.load(path)
            data["semantic_styles"]["salutation"] = data["semantic_styles"]["identity"]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "document_hierarchy"), "fail", entry)
        elif mutation == "duplicate_section_purpose":
            path = stages / "resume.json"
            data = self.load(path)
            data["sections"].append({"id": "different-id", "heading": data["sections"][0]["heading"]})
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "persuasive_content"), "fail", entry)
        elif mutation == "wrong_ledger_location":
            path = stages / "coverage-ledger.json"
            data = self.load(path)
            next(item for item in data["entries"] if item["claim_id"] == "platform-claim")["final_locations"] = ["cover_letter"]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "substance_preservation"), "fail", entry)
        elif mutation == "stale_reviewed_input_hash":
            path = stages / "review-record.json"
            data = self.load(path)
            data["reviewed_input_hashes"]["resume.json"] = "0" * 64
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "exact_artifact_review"), "fail", entry)
        elif mutation == "accepted_region_missing_description":
            path = stages / "change-set.json"
            data = self.load(path)
            page_hash = self.load(stages / "render-evidence.json")["artifacts"][0]["pages"][0]["sha256"]
            data["accepted_regions"] = [{"artifact_id": "resume_page_1", "sha256": page_hash, "description": ""}]
            self.save(path, data)
            with self.assertRaises(ContractError, msg=entry):
                run_manifest(self.manifest)
        elif mutation == "missing_reviewed_requirement_hash":
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 0, report)
            path = self.run_dir / "audit" / "review-record.json"
            data = self.load(path)
            data["reviewed_input_hashes"].pop("opportunity-brief.json")
            self.save(path, data)
            report = validate_run(self.run_dir)
            self.assertEqual(self.gate_status(report, "exact_artifact_review"), "fail", entry)
        elif mutation == "missing_primary_source":
            path = stages / "evidence-bundle.json"
            data = self.load(path)
            next(source for source in data["sources"] if source["kind"] == "baseline").pop("sha256")
            self.save(path, data)
            with self.assertRaises(ContractError, msg=entry):
                run_manifest(self.manifest)
        elif mutation == "conflicting_provenance_identity":
            path = stages / "evidence-bundle.json"
            data = self.load(path)
            data["provenance_requirements"].append({
                "child_entity_id": "client_acme", "parent_entity_id": "independent"
            })
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "career_provenance"), "fail", entry)
        elif mutation == "unrelated_render_page":
            page = stages / "render-pages" / "resume_pdf" / "page-1.png"
            Image.new("RGB", (120, 160), (0, 0, 0)).save(page)
            evidence_path = stages / "render-evidence.json"
            evidence = self.load(evidence_path)
            record = next(item for item in evidence["artifacts"] if item["id"] == "resume_pdf")["pages"][0]
            record["sha256"] = sha256_file(page)
            record["bytes"] = page.stat().st_size
            self.save(evidence_path, evidence)
            review_path = stages / "review-record.json"
            review = self.load(review_path)
            review["page_image_hashes"]["resume_pdf"][0] = record["sha256"]
            self.save(review_path, review)
            with self.assertRaisesRegex(ContractError, "does not match a fresh render", msg=entry):
                run_manifest(self.manifest)
        elif mutation == "missing_user_correction_source":
            path = stages / "evidence-bundle.json"
            data = self.load(path)
            data["claims"][0]["source_ids"] = ["direct-user-correction"]
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "evidence_integrity"), "fail", entry)
        elif mutation == "missing_docx_review_hash":
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 0, report)
            path = self.run_dir / "audit" / "review-record.json"
            data = self.load(path)
            data["artifact_hashes"].pop("resume_docx")
            self.save(path, data)
            report = validate_run(self.run_dir)
            self.assertEqual(self.gate_status(report, "exact_artifact_review"), "fail", entry)
        elif mutation == "baseline_without_hash":
            path = stages / "evidence-bundle.json"
            data = self.load(path)
            next(source for source in data["sources"] if source["kind"] == "baseline").pop("sha256")
            self.save(path, data)
            with self.assertRaises(ContractError, msg=entry):
                run_manifest(self.manifest)
        elif mutation == "missing_visual_correction_evidence":
            path = stages / "change-set.json"
            data = self.load(path)
            data["changes"] = [{
                "id": "change-one", "target_ids": ["identity"],
                "invalidates": ["render", "review", "release"],
            }]
            data["visual_evidence"] = {"reviewed_change_ids": [], "after_page_hashes": []}
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "change_isolation"), "fail", entry)
        elif mutation == "invalid_docx_metadata":
            docx_path = self.root / "artifacts" / "resume.docx"
            document = Document()
            for line in ("Avery Candidate", "Platform protected achievement", "Independent research"):
                document.add_paragraph(line)
            document.save(docx_path)
            evidence_path = stages / "render-evidence.json"
            evidence = self.load(evidence_path)
            item = next(item for item in evidence["artifacts"] if item["id"] == "resume_docx")
            item["sha256"] = sha256_file(docx_path)
            item["bytes"] = docx_path.stat().st_size
            rendered = render_exact_artifact(
                docx_path,
                "docx",
                evidence["tool_capabilities"],
                stages / "render-pages" / "resume_docx",
            )
            item["pages"] = [
                {
                    "page": page_number,
                    "path": str(path.relative_to(stages)),
                    "sha256": sha256_file(path),
                    "bytes": path.stat().st_size,
                }
                for page_number, path in enumerate(rendered, start=1)
            ]
            self.save(evidence_path, evidence)
            review_path = stages / "review-record.json"
            review = self.load(review_path)
            review["artifact_hashes"]["resume_docx"] = item["sha256"]
            review["page_image_hashes"]["resume_docx"] = [page["sha256"] for page in item["pages"]]
            self.save(review_path, review)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, report)
            self.assertEqual(self.gate_status(report, "layout_quality"), "fail", entry)
            errors = next(gate["errors"] for gate in report["gates"] if gate["id"] == "layout_quality")
            self.assertTrue(any("python-docx" in error or "thumbnail" in error for error in errors), errors)
        elif mutation == "dangling_atomic_separator":
            path = self.root / "dangling-degree.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Northbridge University | M.S. & B.S., Electrical &")
            paragraph.add_run().add_break()
            paragraph.add_run("Computer Engineering")
            document.save(path)
            findings = atomic_phrase_findings(
                path,
                "docx",
                ["Northbridge University | M.S. & B.S., Electrical & Computer Engineering"],
            )
            self.assertTrue(any("dangling separator" in finding for finding in findings), findings)
        elif mutation == "fabricated_review_without_transcript":
            path = stages / "review-evidence.json"
            data = self.load(path)
            data.pop("exchanges")
            data["verdict"] = "pass"
            data["unresolved_findings"] = []
            self.save(path, data)
            review_path = stages / "review-record.json"
            review = self.load(review_path)
            review["findings"] = []
            review["unresolved_findings"] = []
            review["external_review_evidence"]["sha256"] = sha256_file(path)
            self.save(review_path, review)
            with self.assertRaisesRegex(ContractError, "exchanges.*required", msg=entry):
                run_manifest(self.manifest)
        elif mutation == "unrelated_pass_transcript_with_current_hashes":
            path = stages / "review-evidence.json"
            data = self.load(path)
            for exchange in data["exchanges"]:
                request = "Approve an unrelated draft. The exact review packet is not identified."
                exchange["request"] = request
                exchange["request_sha256"] = sha256_bytes(request.encode("utf-8"))
                exchange["verdict"] = "pass"
            data["verdict"] = "pass"
            data["unresolved_findings"] = []
            self.save(path, data)
            review_path = stages / "review-record.json"
            review = self.load(review_path)
            review["external_review_evidence"]["sha256"] = sha256_file(path)
            self.save(review_path, review)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "exact_artifact_review"), "fail", entry)
            errors = next(
                gate["errors"]
                for gate in report["gates"]
                if gate["id"] == "exact_artifact_review"
            )
            self.assertTrue(
                any("does not identify the current review packet" in error for error in errors),
                errors,
            )
        elif mutation in {
            "tampered_derived_link_inventory",
            "tampered_derived_emphasis",
            "missing_accessible_document_title",
        }:
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 0, report)
            path = self.run_dir / "audit" / "render-manifest.json"
            data = self.load(path)
            if mutation == "tampered_derived_link_inventory":
                item = next(item for item in data["artifacts"] if item["id"] == "resume_docx")
                item["external_links"] = ["https://unrelated.example.test/profile"]
                expected = "external links differ"
            elif mutation == "tampered_derived_emphasis":
                item = next(item for item in data["artifacts"] if item["id"] == "cover_letter_docx")
                item["emphasis_sha256"] = "0" * 64
                expected = "semantic emphasis differs"
            else:
                item = next(item for item in data["artifacts"] if item["id"] == "resume_pdf")
                item["accessibility"]["document_title"] = ""
                expected = "accessible document title"
            self.save(path, data)
            report = validate_run(self.run_dir)
            self.assertEqual(self.gate_status(report, "cross_format_parity"), "fail", entry)
            errors = next(
                gate["errors"]
                for gate in report["gates"]
                if gate["id"] == "cross_format_parity"
            )
            self.assertTrue(any(expected in error for error in errors), errors)
        elif mutation == "unresolvable_render_artifact_source":
            path = stages / "render-evidence.json"
            data = self.load(path)
            item = next(item for item in data["artifacts"] if item["id"] == "resume_pdf")
            item["source"] = str(self.root / "artifacts" / "cover-letter.pdf")
            self.save(path, data)
            with self.assertRaisesRegex(ContractError, "Render evidence source for resume_pdf", msg=entry):
                run_manifest(self.manifest)
        elif mutation == "stage_source_byte_rewrite":
            source = stages / "strategy.json"
            value = self.load(source)
            source.write_text(
                json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n",
                encoding="utf-8",
            )
            evidence_path = stages / "review-evidence.json"
            evidence = self.load(evidence_path)
            evidence["reviewed_input_hashes"]["strategy.json"] = sha256_file(source)
            packet_hash = sha256_bytes(canonical_json({
                "schema_version": "1.0",
                "artifact_hashes": evidence["artifact_hashes"],
                "page_image_hashes": evidence["page_image_hashes"],
                "reviewed_input_hashes": evidence["reviewed_input_hashes"],
            }))
            old_packet_hash = evidence["review_packet_sha256"]
            evidence["review_packet_sha256"] = packet_hash
            for exchange in evidence["exchanges"]:
                exchange["review_packet_sha256"] = packet_hash
                exchange["request"] = exchange["request"].replace(old_packet_hash, packet_hash)
                exchange["request_sha256"] = sha256_bytes(exchange["request"].encode("utf-8"))
            self.save(evidence_path, evidence)
            review_path = stages / "review-record.json"
            review = self.load(review_path)
            review["reviewed_input_hashes"]["strategy.json"] = sha256_file(source)
            review["external_review_evidence"]["sha256"] = sha256_file(evidence_path)
            self.save(review_path, review)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 0, report)
            copied = self.run_dir / "audit" / "strategy.json"
            self.assertEqual(copied.read_bytes(), source.read_bytes())
        elif mutation == "unsupported_semantic_assertion":
            path = stages / "resume.json"
            data = self.load(path)
            data["semantic_text"] += "\nReceived an unsupported global award."
            self.save(path, data)
            code, report = run_manifest(self.manifest)
            self.assertEqual(code, 2, entry)
            self.assertEqual(self.gate_status(report, "evidence_integrity"), "fail", entry)
        else:
            self.fail(f"Unknown failure mutation {mutation!r}: {entry}")


def _make_failure_test(entry: dict):
    def test(self: FailureMutationTests) -> None:
        self.run_failure_mutation(entry)

    test.__name__ = entry["automated_test"]
    test.__doc__ = f"Regression {entry['id']}: {entry['guard']}"
    return test


for _entry in CATALOG:
    setattr(FailureMutationTests, _entry["automated_test"], _make_failure_test(_entry))


class FailureCatalogBindingTests(unittest.TestCase):
    def test_every_catalog_entry_resolves_to_one_executable_mutation_test(self) -> None:
        self.assertEqual([entry["id"] for entry in CATALOG], list(range(1, len(CATALOG) + 1)))
        test_names = [entry["automated_test"] for entry in CATALOG]
        self.assertEqual(len(test_names), len(set(test_names)))
        for entry in CATALOG:
            self.assertTrue(hasattr(FailureMutationTests, entry["automated_test"]), entry)
            self.assertNotIn("manual_evidence_gate", entry)


if __name__ == "__main__":
    unittest.main()
