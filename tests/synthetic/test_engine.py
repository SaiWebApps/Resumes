from __future__ import annotations

import json
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen.canvas import Canvas

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "src"))

from application_pack.artifacts import canonical_json, extract_docx_text, normalize_semantic_text, sha256_bytes, sha256_file  # noqa: E402
from application_pack.contracts import ContractError  # noqa: E402
from application_pack.engine import run_manifest, validate_run  # noqa: E402
from application_pack.rendering import render_exact_artifact  # noqa: E402
from application_pack.validation import review_packet_sha256  # noqa: E402


def write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def make_pdf(path: Path, pages: list[list[str]]) -> None:
    canvas = Canvas(str(path), pagesize=letter)
    for page in pages:
        y = 740
        for line in page:
            canvas.drawString(72, y, line)
            y -= 18
        canvas.showPage()
    canvas.save()


def make_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(path)
    _sanitize_synthetic_docx(path, lines)


def _sanitize_synthetic_docx(path: Path, lines: list[str]) -> None:
    pages = 2 if "resume" in path.name else 1
    words = max(1, sum(len(line.split()) for line in lines))
    characters = max(1, sum(len(line) for line in lines))
    characters_with_spaces = characters + max(0, len(lines) - 1)
    core = f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"><dc:title>Synthetic application document</dc:title><dc:subject>Regression fixture</dc:subject><dc:creator>Avery Candidate</dc:creator><cp:keywords/><dc:description>Portable application-package regression fixture</dc:description><cp:lastModifiedBy>Avery Candidate</cp:lastModifiedBy><cp:revision>1</cp:revision><dcterms:created xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:created><dcterms:modified xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified><cp:category/></cp:coreProperties>""".encode()
    app = f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Pages>{pages}</Pages><Words>{words}</Words><Characters>{characters}</Characters><CharactersWithSpaces>{characters_with_spaces}</CharactersWithSpaces></Properties>""".encode()
    replacement = path.with_suffix(".clean.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename == "docProps/thumbnail.jpeg":
                continue
            payload = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                payload = core
            elif info.filename == "docProps/app.xml":
                payload = app
            elif info.filename == "word/settings.xml":
                payload = re.sub(rb"<w:savePreviewPicture(?:\s[^>]*)?\s*/>", b"", payload)
            elif info.filename == "_rels/.rels":
                relationships = ElementTree.fromstring(payload)
                for relationship in list(relationships):
                    if relationship.get("Target", "").endswith("docProps/thumbnail.jpeg"):
                        relationships.remove(relationship)
                payload = ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True)
            target.writestr(info, payload)
    replacement.replace(path)


def make_synthetic_renderers(root: Path) -> dict:
    renderers = root / "renderers"
    renderers.mkdir(parents=True)
    pdftoppm = renderers / "synthetic-pdftoppm.py"
    pdftoppm.write_text(
        f"""#!{sys.executable}
import binascii
import hashlib
import re
import struct
import sys
import zlib
from pathlib import Path

def write_png(path, color):
    width, height = 120, 160
    raw = b"".join(b"\\x00" + bytes(color) * width for _ in range(height))
    def chunk(kind, payload):
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", binascii.crc32(kind + payload) & 0xffffffff)
    path.write_bytes(
        b"\\x89PNG\\r\\n\\x1a\\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )

source = Path(sys.argv[-2])
prefix = Path(sys.argv[-1])
payload = source.read_bytes()
page_count = len(re.findall(rb"/Type\\s*/Page(?!s)", payload)) or 1
for page in range(1, page_count + 1):
    digest = hashlib.sha256(payload + str(page).encode()).digest()
    write_png(Path(f"{{prefix}}-{{page}}.png"), tuple(digest[:3]))
""",
        encoding="utf-8",
    )
    pdftoppm.chmod(0o755)

    docx_renderer = renderers / "synthetic-docx-renderer.py"
    docx_renderer.write_text(
        """import argparse
import binascii
import hashlib
import struct
import zlib
from pathlib import Path

def write_png(path, color):
    width, height = 120, 160
    raw = b"".join(b"\\x00" + bytes(color) * width for _ in range(height))
    def chunk(kind, payload):
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", binascii.crc32(kind + payload) & 0xffffffff)
    path.write_bytes(
        b"\\x89PNG\\r\\n\\x1a\\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )

parser = argparse.ArgumentParser()
parser.add_argument("source")
parser.add_argument("--output_dir", required=True)
parser.add_argument("--emit_pdf", action="store_true")
args = parser.parse_args()
source = Path(args.source)
output = Path(args.output_dir)
output.mkdir(parents=True, exist_ok=True)
page_count = 2 if "resume" in source.name else 1
payload = source.read_bytes()
for page in range(1, page_count + 1):
    digest = hashlib.sha256(payload + str(page).encode()).digest()
    write_png(output / f"page-{page}.png", tuple(digest[:3]))
""",
        encoding="utf-8",
    )
    return {
        "pdftoppm": {"path": str(pdftoppm), "sha256": sha256_file(pdftoppm)},
        "docx_renderer": {"path": str(docx_renderer), "sha256": sha256_file(docx_renderer)},
        "python": {"path": str(Path(sys.executable).absolute()), "sha256": sha256_file(Path(sys.executable).absolute())},
    }


def make_fixture(root: Path) -> Path:
    stages = root / "stages"
    artifacts = root / "artifacts"
    stages.mkdir(parents=True)
    artifacts.mkdir(parents=True)

    resume_lines = ["Avery Candidate", "Platform protected achievement", "Independent research"]
    cover_lines = ["Dear Program Team", "Consulting client achievement", "Sincerely", "Avery Candidate"]
    make_pdf(artifacts / "resume.pdf", [resume_lines[:2], resume_lines[2:]])
    make_docx(artifacts / "resume.docx", resume_lines)
    make_pdf(artifacts / "cover-letter.pdf", [cover_lines])
    make_docx(artifacts / "cover-letter.docx", cover_lines)

    evidence = {
        "schema_version": "1.0",
        "profile_id": "synthetic",
        "sources": [{"id": "source-baseline", "kind": "baseline", "uri": "fixture://baseline", "sha256": "0" * 64}],
        "claims": [
            {"id": "identity-claim", "text": "Avery Candidate", "source_ids": ["source-baseline"], "protected": True, "entity_id": "candidate", "parent_entity_id": None, "required_in": ["resume", "cover_letter"]},
            {"id": "platform-claim", "text": "Platform protected achievement", "source_ids": ["source-baseline"], "protected": True, "entity_id": "primary_employer", "parent_entity_id": "primary_employer", "required_in": ["resume"]},
            {"id": "research-claim", "text": "Independent research", "source_ids": ["source-baseline"], "protected": True, "entity_id": "research", "parent_entity_id": "independent", "required_in": ["resume"]},
            {"id": "client-claim", "text": "Consulting client achievement", "source_ids": ["source-baseline"], "protected": True, "entity_id": "client_acme", "parent_entity_id": "consultancy", "required_in": ["cover_letter"]}
        ],
        "provenance_requirements": [{"child_entity_id": "client_acme", "parent_entity_id": "consultancy"}],
        "conflicts": []
    }
    resume = {
        "schema_version": "1.0", "document_id": "resume", "semantic_text": "\n".join(resume_lines),
        "claims": [
            {"claim_id": "identity-claim", "text": "Avery Candidate", "section_id": "identity", "entity_id": "candidate", "parent_entity_id": None},
            {"claim_id": "platform-claim", "text": "Platform protected achievement", "section_id": "experience", "entity_id": "primary_employer", "parent_entity_id": "primary_employer"},
            {"claim_id": "research-claim", "text": "Independent research", "section_id": "research", "entity_id": "research", "parent_entity_id": "independent"}
        ],
        "sections": [{"id": "experience", "heading": "Experience"}, {"id": "research", "heading": "Research"}],
        "semantic_units": [
            {"id": "resume-line-1", "text": "Avery Candidate", "kind": "evidence", "support_refs": ["claim:identity-claim"]},
            {"id": "resume-line-2", "text": "Platform protected achievement", "kind": "evidence", "support_refs": ["claim:platform-claim"]},
            {"id": "resume-line-3", "text": "Independent research", "kind": "evidence", "support_refs": ["claim:research-claim"]}
        ]
    }
    cover = {
        "schema_version": "1.0", "document_id": "cover_letter", "semantic_text": "\n".join(cover_lines),
        "claims": [
            {"claim_id": "client-claim", "text": "Consulting client achievement", "section_id": "credibility", "entity_id": "client_acme", "parent_entity_id": "consultancy"},
            {"claim_id": "identity-claim", "text": "Avery Candidate", "section_id": "signature", "entity_id": "candidate", "parent_entity_id": None}
        ],
        "proposal_ids": ["topic-1"],
        "semantic_units": [
            {"id": "cover-line-1", "text": "Dear Program Team", "kind": "structure", "support_refs": ["structure:salutation", "opportunity:https://example.test/conference"]},
            {"id": "cover-line-2", "text": "Consulting client achievement", "kind": "evidence", "support_refs": ["claim:client-claim"]},
            {"id": "cover-line-3", "text": "Sincerely", "kind": "structure", "support_refs": ["structure:closing"]},
            {"id": "cover-line-4", "text": "Avery Candidate", "kind": "evidence", "support_refs": ["claim:identity-claim"]}
        ]
    }
    coverage = {
        "schema_version": "1.0",
        "entries": [
            {"claim_id": "identity-claim", "status": "retained", "final_locations": ["resume", "cover_letter"], "strength_preserved": True},
            {"claim_id": "platform-claim", "status": "retained", "final_locations": ["resume"], "strength_preserved": True},
            {"claim_id": "research-claim", "status": "retained", "final_locations": ["resume"], "strength_preserved": True},
            {"claim_id": "client-claim", "status": "retained", "final_locations": ["cover_letter"], "strength_preserved": True}
        ]
    }
    opportunity = {
        "schema_version": "1.0", "opportunity_url": "https://example.test/conference",
        "sources": [{"url": "https://example.test/conference", "title": "Conference"}],
        "audiences": ["city leaders"], "selection_criteria": ["evidence-backed expertise"],
        "tracks": ["AI"], "submission_requirements": ["resume", "cover letter"]
    }
    strategy = {
        "schema_version": "1.0", "positioning": "Production AI leader and educator",
        "credibility_arc": ["engineering", "teaching"],
        "proposed_topics": [{"id": "topic-1", "title": "Trustworthy AI", "evidence_claim_ids": ["platform-claim"], "attendee_outcomes": ["Choose evidence before scale"]}],
        "resume_purpose": "Establish the complete record", "cover_letter_purpose": "Connect evidence to the audience"
    }
    design = {
        "schema_version": "1.0",
        "visible_elements": [
            {"document_id": "resume", "element_id": "identity", "function": "identity", "reader_purpose": "Identify the candidate"},
            {"document_id": "cover_letter", "element_id": "salutation", "function": "salutation", "reader_purpose": "Address the program team"}
        ],
        "semantic_styles": {"identity": "Name", "salutation": "Salutation"},
        "page_constraints": {"resume": 2, "cover_letter": 1}
    }
    changes = {
        "schema_version": "1.0",
        "changes": [],
        "accepted_regions": [],
        "collateral_change_ids": [],
        "visual_evidence": {"reviewed_change_ids": [], "after_page_hashes": []},
    }

    write_json(stages / "evidence-bundle.json", evidence)
    write_json(stages / "opportunity-brief.json", opportunity)
    write_json(stages / "strategy.json", strategy)
    write_json(stages / "resume.json", resume)
    write_json(stages / "coverage-ledger.json", coverage)
    write_json(stages / "cover-letter.json", cover)
    write_json(stages / "design-contract.json", design)
    write_json(stages / "change-set.json", changes)

    artifact_hashes = {
        "resume_pdf": sha256_file(artifacts / "resume.pdf"),
        "resume_docx": sha256_file(artifacts / "resume.docx"),
        "cover_letter_pdf": sha256_file(artifacts / "cover-letter.pdf"),
        "cover_letter_docx": sha256_file(artifacts / "cover-letter.docx"),
    }
    renderer_capabilities = make_synthetic_renderers(root)
    render_pages: dict[str, list[dict]] = {}
    artifact_paths = {
        "resume_pdf": artifacts / "resume.pdf",
        "resume_docx": artifacts / "resume.docx",
        "cover_letter_pdf": artifacts / "cover-letter.pdf",
        "cover_letter_docx": artifacts / "cover-letter.docx",
    }
    render_counts = {"resume_pdf": 2, "resume_docx": 2, "cover_letter_pdf": 1, "cover_letter_docx": 1}
    render_records = []
    for artifact_id, count in render_counts.items():
        output_dir = stages / "render-pages" / artifact_id
        generated_pages = render_exact_artifact(
            artifact_paths[artifact_id],
            artifact_paths[artifact_id].suffix.lstrip("."),
            renderer_capabilities,
            output_dir,
        )
        if len(generated_pages) != count:
            raise AssertionError(f"synthetic renderer produced {len(generated_pages)} pages for {artifact_id}")
        pages = []
        for page_number, page_path in enumerate(generated_pages, start=1):
            pages.append({
                "page": page_number,
                "path": str(page_path.relative_to(stages)),
                "sha256": sha256_file(page_path),
                "bytes": page_path.stat().st_size,
            })
        render_pages[artifact_id] = pages
        source = artifact_paths[artifact_id]
        render_records.append({
            "id": artifact_id,
            "source": str(source),
            "format": source.suffix.lstrip("."),
            "sha256": artifact_hashes[artifact_id],
            "bytes": source.stat().st_size,
            "pages": pages,
        })
    write_json(stages / "render-evidence.json", {
        "schema_version": "1.0",
        "tool_capabilities": renderer_capabilities,
        "artifacts": render_records,
    })
    reviewed_input_hashes = {
        filename: sha256_file(stages / filename)
        for filename in (
            "evidence-bundle.json", "opportunity-brief.json", "strategy.json", "resume.json",
            "coverage-ledger.json", "cover-letter.json", "design-contract.json", "change-set.json",
        )
    }
    page_image_hashes = {
        artifact_id: [page["sha256"] for page in pages]
        for artifact_id, pages in render_pages.items()
    }
    review_packet_sha256 = sha256_bytes(canonical_json({
        "schema_version": "1.0",
        "artifact_hashes": artifact_hashes,
        "page_image_hashes": page_image_hashes,
        "reviewed_input_hashes": reviewed_input_hashes,
    }))
    primary_request = json.dumps(
        {
            "instruction": "Review the exact synthetic artifacts and pages against the supplied requirements.",
            "review_packet_sha256": review_packet_sha256,
            "artifact_hashes": artifact_hashes,
            "page_image_hashes": page_image_hashes,
            "reviewed_input_hashes": reviewed_input_hashes,
        },
        sort_keys=True,
    )
    primary_response = "PASS: exact artifacts, pages, evidence, and requirements have no unresolved findings."
    reciprocal_request = json.dumps(
        {
            "instruction": "Audit the independent review and exact hash packet for unsupported approval.",
            "review_packet_sha256": review_packet_sha256,
            "primary_response_sha256": sha256_bytes(primary_response.encode("utf-8")),
            "artifact_hashes": artifact_hashes,
            "page_image_hashes": page_image_hashes,
        },
        sort_keys=True,
    )
    reciprocal_response = "PASS: the external review is hash-bound, reciprocal, and supported by the current packet."
    write_json(stages / "review-evidence.json", {
        "schema_version": "1.0",
        "evidence_id": "synthetic-external-review",
        "review_mode": "cross_model",
        "primary_reviewer": {"provider": "synthetic-independent-reviewer", "model": "synthetic-review-model"},
        "artifact_hashes": artifact_hashes,
        "page_image_hashes": page_image_hashes,
        "reviewed_input_hashes": reviewed_input_hashes,
        "review_packet_sha256": review_packet_sha256,
        "exchanges": [
            {
                "requester_provider": "synthetic-codex",
                "reviewer_provider": "synthetic-independent-reviewer",
                "reviewer_model": "synthetic-review-model",
                "review_packet_sha256": review_packet_sha256,
                "request": primary_request,
                "request_sha256": sha256_bytes(primary_request.encode("utf-8")),
                "response": primary_response,
                "response_sha256": sha256_bytes(primary_response.encode("utf-8")),
                "verdict": "pass",
            },
            {
                "requester_provider": "synthetic-independent-reviewer",
                "reviewer_provider": "synthetic-codex",
                "reviewer_model": "synthetic-codex-model",
                "review_packet_sha256": review_packet_sha256,
                "request": reciprocal_request,
                "request_sha256": sha256_bytes(reciprocal_request.encode("utf-8")),
                "response": reciprocal_response,
                "response_sha256": sha256_bytes(reciprocal_response.encode("utf-8")),
                "verdict": "pass",
            },
        ],
        "reciprocal_review": {
            "completed": True,
            "primary_exchange_index": 0,
            "reciprocal_exchange_index": 1,
        },
        "verdict": "pass",
        "findings": [],
        "unresolved_findings": [],
    })
    write_json(stages / "review-record.json", {
        "schema_version": "1.0", "mode": "cross_model", "provider": "synthetic-independent-reviewer",
        "artifact_hashes": artifact_hashes, "findings": [], "unresolved_findings": [],
        "page_image_hashes": page_image_hashes,
        "reviewed_input_hashes": reviewed_input_hashes,
        "external_review_evidence": {
            "path": "review-evidence.json",
            "sha256": sha256_file(stages / "review-evidence.json"),
        },
        "manual_checks": {
            "full_page_readability": True,
            "physical_edge_clearance": True,
            "natural_content_flow": True,
            "continuous_page_review": True,
            "metric_visibility": True,
            "talk_taxonomy": True,
            "current_link_inventory": True,
            "docx_hanging_indents": True,
            "role_date_rows_atomic": True,
            "source_emphasis_preservation": True,
        },
        "visual_review": {"all_pages_inspected": True, "blocking_defects": []}
    })

    manifest = {
        "schema_version": "1.0", "run_id": "synthetic-run", "profile": {"id": "synthetic"},
        "objective": "Produce a verified conference speaker package", "audience": "conference program team",
        "opportunity": {"url": "https://example.test/conference", "referral_context": "Referral"},
        "output_dir": "runs",
        "stage_sources": {
            "evidence_bundle": "stages/evidence-bundle.json", "opportunity_brief": "stages/opportunity-brief.json",
            "strategy": "stages/strategy.json", "resume": "stages/resume.json",
            "coverage_ledger": "stages/coverage-ledger.json", "cover_letter": "stages/cover-letter.json",
            "design_contract": "stages/design-contract.json", "change_set": "stages/change-set.json",
            "render_evidence": "stages/render-evidence.json",
            "review_evidence": "stages/review-evidence.json",
            "review_record": "stages/review-record.json"
        },
        "artifact_sources": {
            "resume_pdf": "artifacts/resume.pdf", "resume_docx": "artifacts/resume.docx",
            "cover_letter_pdf": "artifacts/cover-letter.pdf", "cover_letter_docx": "artifacts/cover-letter.docx"
        },
        "constraints": {"page_counts": {"resume": 2, "cover_letter": 1}, "forbid_page_numbers": True},
        "review_policy": {"allow_disclosed_fallback": True},
        "capabilities": {"filesystem": True, "research": True, "rendering": True, "visual_inspection": True}
    }
    path = root / "run-manifest.json"
    write_json(path, manifest)
    return path


class EngineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.manifest = make_fixture(self.root)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def load(self, path: Path) -> dict:
        return json.loads(path.read_text(encoding="utf-8"))

    def save(self, path: Path, value: dict) -> None:
        write_json(path, value)

    @property
    def run_dir(self) -> Path:
        return self.root / "runs" / "synthetic-run"

    def assert_gate_failed(self, report: dict, gate_id: str) -> None:
        statuses = {gate["id"]: gate["status"] for gate in report["gates"]}
        self.assertEqual(statuses.get(gate_id), "fail", report)

    def test_valid_release_promotes_exactly_four_submission_files(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        self.assertEqual(report["status"], "pass")
        self.assertEqual(report["score"], 100)
        self.assertEqual(sorted(path.name for path in (self.run_dir / "final").iterdir()),
                         ["cover-letter.docx", "cover-letter.pdf", "resume.docx", "resume.pdf"])
        self.assertEqual(validate_run(self.run_dir)["status"], "pass")

    def test_promotion_failure_never_records_released_state(self) -> None:
        with patch("application_pack.engine._promote", side_effect=OSError("synthetic promotion failure")):
            with self.assertRaisesRegex(OSError, "synthetic promotion failure"):
                run_manifest(self.manifest)

        state = self.load(self.run_dir / "work" / "state.json")
        self.assertNotEqual(state["status"], "released")
        self.assertFalse((self.run_dir / "final").exists())

    def test_released_state_without_final_directory_fails_revalidation(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        final = self.run_dir / "final"
        for artifact in final.iterdir():
            artifact.unlink()
        final.rmdir()

        revalidated = validate_run(self.run_dir)
        self.assert_gate_failed(revalidated, "submission_readiness")
        readiness = next(gate for gate in revalidated["gates"] if gate["id"] == "submission_readiness")
        self.assertTrue(any("released but final/ is missing" in error for error in readiness["errors"]), readiness)

    def test_unsupported_claim_blocks_release(self) -> None:
        path = self.root / "stages" / "resume.json"
        data = self.load(path)
        data["claims"].append({"claim_id": "invented", "text": "Invented", "section_id": "experience", "entity_id": "x", "parent_entity_id": "x"})
        self.save(path, data)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "evidence_integrity")
        self.assertFalse((self.run_dir / "final").exists())

    def test_unresolved_material_evidence_conflict_blocks_release(self) -> None:
        path = self.root / "stages" / "evidence-bundle.json"
        data = self.load(path)
        data["conflicts"].append({
            "id": "conflict-employment-scope",
            "description": "Two accepted sources disagree about management scope.",
            "impact": "The disagreement changes a protected leadership claim.",
            "material": True,
            "status": "unresolved",
            "source_ids": ["source-baseline"],
        })
        self.save(path, data)

        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2, report)
        self.assert_gate_failed(report, "evidence_integrity")
        evidence_gate = next(gate for gate in report["gates"] if gate["id"] == "evidence_integrity")
        self.assertTrue(any("Material evidence conflict" in error for error in evidence_gate["errors"]), evidence_gate)
        self.assertFalse((self.run_dir / "final").exists())

    def test_omitted_protected_claim_blocks_release(self) -> None:
        path = self.root / "stages" / "coverage-ledger.json"
        data = self.load(path)
        data["entries"] = [entry for entry in data["entries"] if entry["claim_id"] != "platform-claim"]
        self.save(path, data)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "substance_preservation")

    def test_bad_provenance_blocks_release(self) -> None:
        path = self.root / "stages" / "cover-letter.json"
        data = self.load(path)
        data["claims"][0]["parent_entity_id"] = "independent"
        self.save(path, data)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "career_provenance")

    def test_purposeless_element_blocks_release(self) -> None:
        path = self.root / "stages" / "design-contract.json"
        data = self.load(path)
        data["visible_elements"][0]["reader_purpose"] = ""
        self.save(path, data)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "document_hierarchy")

    def test_page_number_element_blocks_release(self) -> None:
        path = self.root / "stages" / "design-contract.json"
        data = self.load(path)
        data["visible_elements"].append({"document_id": "resume", "element_id": "page-number", "function": "page_number", "reader_purpose": "Show the page"})
        self.save(path, data)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "document_hierarchy")

    def test_stale_review_hash_blocks_release(self) -> None:
        path = self.root / "stages" / "review-record.json"
        data = self.load(path)
        data["artifact_hashes"]["resume_pdf"] = "0" * 64
        self.save(path, data)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "exact_artifact_review")

    def test_stale_review_packet_digest_blocks_release(self) -> None:
        evidence_path = self.root / "stages" / "review-evidence.json"
        evidence = self.load(evidence_path)
        evidence["review_packet_sha256"] = "0" * 64
        self.save(evidence_path, evidence)
        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["external_review_evidence"]["sha256"] = sha256_file(evidence_path)
        self.save(review_path, review)

        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2, report)
        self.assert_gate_failed(report, "exact_artifact_review")
        errors = next(
            gate["errors"]
            for gate in report["gates"]
            if gate["id"] == "exact_artifact_review"
        )
        self.assertTrue(any("review packet digest is stale" in error for error in errors), errors)

    def test_cross_format_drift_blocks_release(self) -> None:
        artifact_path = self.root / "artifacts" / "resume.docx"
        make_docx(artifact_path, ["Different content"])
        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        item = next(item for item in evidence["artifacts"] if item["id"] == "resume_docx")
        item["sha256"] = sha256_file(artifact_path)
        item["bytes"] = artifact_path.stat().st_size
        rendered = render_exact_artifact(
            artifact_path,
            "docx",
            evidence["tool_capabilities"],
            self.root / "stages" / "render-pages" / "resume_docx",
        )
        item["pages"] = [
            {
                "page": page_number,
                "path": str(path.relative_to(self.root / "stages")),
                "sha256": sha256_file(path),
                "bytes": path.stat().st_size,
            }
            for page_number, path in enumerate(rendered, start=1)
        ]
        self.save(evidence_path, evidence)
        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["artifact_hashes"]["resume_docx"] = item["sha256"]
        review["page_image_hashes"]["resume_docx"] = [page["sha256"] for page in item["pages"]]
        self.save(review_path, review)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2)
        self.assert_gate_failed(report, "cross_format_parity")

    def test_cross_format_link_inventory_drift_blocks_validation(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        path = self.run_dir / "audit" / "render-manifest.json"
        data = self.load(path)
        item = next(item for item in data["artifacts"] if item["id"] == "resume_docx")
        item["external_links"] = ["https://different.example.test/profile"]
        self.save(path, data)

        report = validate_run(self.run_dir)
        self.assert_gate_failed(report, "cross_format_parity")
        parity = next(gate for gate in report["gates"] if gate["id"] == "cross_format_parity")
        self.assertTrue(any("external links differ" in error for error in parity["errors"]), parity)

    def test_cross_format_emphasis_drift_blocks_validation(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        path = self.run_dir / "audit" / "render-manifest.json"
        data = self.load(path)
        item = next(item for item in data["artifacts"] if item["id"] == "cover_letter_docx")
        item["emphasis_sha256"] = "0" * 64
        self.save(path, data)

        report = validate_run(self.run_dir)
        self.assert_gate_failed(report, "cross_format_parity")
        parity = next(gate for gate in report["gates"] if gate["id"] == "cross_format_parity")
        self.assertTrue(any("semantic emphasis differs" in error for error in parity["errors"]), parity)

    def test_missing_derived_accessibility_evidence_blocks_validation(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        path = self.run_dir / "audit" / "render-manifest.json"
        data = self.load(path)
        item = next(item for item in data["artifacts"] if item["id"] == "resume_pdf")
        item["accessibility"]["document_title"] = ""
        self.save(path, data)

        report = validate_run(self.run_dir)
        self.assert_gate_failed(report, "cross_format_parity")
        parity = next(gate for gate in report["gates"] if gate["id"] == "cross_format_parity")
        self.assertTrue(any("accessible document title" in error for error in parity["errors"]), parity)

    def test_no_progress_detector_blocks_second_identical_failure(self) -> None:
        path = self.root / "stages" / "design-contract.json"
        data = self.load(path)
        data["visible_elements"][0]["reader_purpose"] = ""
        self.save(path, data)
        first_code, _ = run_manifest(self.manifest)
        second_code, second = run_manifest(self.manifest)
        self.assertEqual(first_code, 2)
        self.assertEqual(second_code, 3)
        self.assertEqual(second["run_status"], "blocked_no_progress")

    def test_missing_stage_yields_next_action_without_final(self) -> None:
        (self.root / "stages" / "strategy.json").unlink()
        code, result = run_manifest(self.manifest)
        self.assertEqual(code, 4)
        self.assertEqual(result["stage"], "positioning")
        self.assertEqual(result["skill"], "application-positioner")
        self.assertTrue((self.run_dir / "next-action.json").is_file())
        self.assertFalse((self.run_dir / "final").exists())

    def test_post_run_audit_mutation_is_detected_as_stale(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        path = self.run_dir / "audit" / "strategy.json"
        data = self.load(path)
        data["positioning"] = "Changed after release"
        self.save(path, data)
        report = validate_run(self.run_dir)
        self.assert_gate_failed(report, "state_freshness")

    def test_stage_json_is_validated_against_public_schema(self) -> None:
        path = self.root / "stages" / "opportunity-brief.json"
        data = self.load(path)
        data.pop("tracks")
        self.save(path, data)
        with self.assertRaisesRegex(ContractError, r"tracks.*required"):
            run_manifest(self.manifest)
        self.assertFalse((self.run_dir / "final").exists())

    def test_validated_stage_sources_preserve_exact_reviewed_bytes(self) -> None:
        source = self.root / "stages" / "strategy.json"
        value = self.load(source)
        source.write_text(
            json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n",
            encoding="utf-8",
        )
        source_hash = sha256_file(source)

        evidence_path = self.root / "stages" / "review-evidence.json"
        evidence = self.load(evidence_path)
        evidence["reviewed_input_hashes"]["strategy.json"] = source_hash
        packet_hash = review_packet_sha256(
            evidence["artifact_hashes"],
            evidence["page_image_hashes"],
            evidence["reviewed_input_hashes"],
        )
        evidence["review_packet_sha256"] = packet_hash
        for exchange in evidence["exchanges"]:
            request = json.loads(exchange["request"])
            request["review_packet_sha256"] = packet_hash
            if "reviewed_input_hashes" in request:
                request["reviewed_input_hashes"] = evidence["reviewed_input_hashes"]
            exchange["review_packet_sha256"] = packet_hash
            exchange["request"] = json.dumps(request, sort_keys=True)
            exchange["request_sha256"] = sha256_bytes(exchange["request"].encode("utf-8"))
        self.save(evidence_path, evidence)

        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["reviewed_input_hashes"]["strategy.json"] = source_hash
        review["external_review_evidence"]["sha256"] = sha256_file(evidence_path)
        self.save(review_path, review)

        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        copied = self.run_dir / "audit" / "strategy.json"
        self.assertEqual(sha256_file(copied), source_hash)
        self.assertEqual(copied.read_bytes(), source.read_bytes())

    def test_docx_tabs_are_semantic_word_boundaries(self) -> None:
        path = self.root / "artifacts" / "tabbed.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Principal Engineer")
        paragraph.add_run().add_tab()
        paragraph.add_run("Feb 2021")
        document.save(path)
        self.assertEqual(normalize_semantic_text(extract_docx_text(path)), "principal engineer feb 2021")

    def test_wrong_page_count_blocks_release(self) -> None:
        manifest = self.load(self.manifest)
        manifest["constraints"]["page_counts"]["resume"] = 3
        self.save(self.manifest, manifest)
        with self.assertRaisesRegex(ContractError, "resume_pdf has 2 pages; expected 3"):
            run_manifest(self.manifest)

    def test_render_evidence_source_must_resolve_to_exact_artifact(self) -> None:
        path = self.root / "stages" / "render-evidence.json"
        data = self.load(path)
        resume = next(item for item in data["artifacts"] if item["id"] == "resume_pdf")
        resume["source"] = str(self.root / "artifacts" / "cover-letter.pdf")
        self.save(path, data)
        with self.assertRaisesRegex(ContractError, "Render evidence source for resume_pdf"):
            run_manifest(self.manifest)

    def test_failure_catalog_covers_every_failure_with_one_owner_and_executable_mutation(self) -> None:
        catalog = json.loads((Path(__file__).with_name("failure_catalog.json")).read_text(encoding="utf-8"))
        self.assertEqual([entry["id"] for entry in catalog], list(range(1, len(catalog) + 1)))
        owners = {
            "application-package", "career-evidence", "opportunity-auditor", "application-positioner",
            "resume-author", "speaker-pitch-author", "document-compositor",
            "artifact-change-controller", "submission-release-auditor"
        }
        automated_tests: set[str] = set()
        for entry in catalog:
            self.assertIn(entry["primary_owner"], owners, entry)
            self.assertIsInstance(entry["primary_owner"], str, entry)
            automated = entry.get("automated_test")
            self.assertIsInstance(automated, str, entry)
            self.assertTrue(automated.startswith(f"test_failure_{entry['id']:02d}_"), entry)
            self.assertNotIn(automated, automated_tests, entry)
            automated_tests.add(automated)
            self.assertNotIn("manual_evidence_gate", entry)
            mutation = entry.get("mutation")
            self.assertIsInstance(mutation, str, entry)
            self.assertTrue(mutation.strip(), entry)


if __name__ == "__main__":
    unittest.main()
