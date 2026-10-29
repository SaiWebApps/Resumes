from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from PIL import Image

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from application_pack.artifacts import sha256_file  # noqa: E402
from application_pack.contracts import ContractError  # noqa: E402
from application_pack.engine import run_manifest, validate_run  # noqa: E402
from application_pack.rendering import render_exact_artifact  # noqa: E402
from test_engine import make_fixture, write_json  # noqa: E402


class VisualEvidenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.manifest = make_fixture(self.root)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def load(self, path: Path) -> dict:
        return json.loads(path.read_text(encoding="utf-8"))

    def save(self, path: Path, value: dict) -> None:
        write_json(path, value)

    @property
    def run_dir(self) -> Path:
        return self.root / "runs" / "synthetic-run"

    def assert_gate_failed(self, report: dict, gate_id: str) -> None:
        statuses = {gate["id"]: gate["status"] for gate in report["gates"]}
        self.assertEqual(statuses.get(gate_id), "fail", report)

    def test_white_72_point_docx_blocks_release_even_with_fresh_artifact_hashes(self) -> None:
        docx_path = self.root / "artifacts" / "resume.docx"
        document = Document(docx_path)
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(72)
        document.save(docx_path)

        artifact_hash = sha256_file(docx_path)
        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        item = next(item for item in evidence["artifacts"] if item["id"] == "resume_docx")
        item["sha256"] = artifact_hash
        item["bytes"] = docx_path.stat().st_size
        rendered = render_exact_artifact(
            docx_path,
            "docx",
            evidence["tool_capabilities"],
            self.root / "stages" / "render-pages" / "resume_docx",
        )
        item["pages"] = [
            {
                "page": page_number,
                "path": str(path.relative_to(self.root / "stages")),
                "sha256": sha256_file(path),
                "bytes": path.stat().st_size,
            }
            for page_number, path in enumerate(rendered, start=1)
        ]
        self.save(evidence_path, evidence)

        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["artifact_hashes"]["resume_docx"] = artifact_hash
        review["page_image_hashes"]["resume_docx"] = [page["sha256"] for page in item["pages"]]
        self.save(review_path, review)

        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2, report)
        self.assert_gate_failed(report, "layout_quality")
        errors = next(gate["errors"] for gate in report["gates"] if gate["id"] == "layout_quality")
        self.assertTrue(any("white text" in error or "72 pt" in error for error in errors), errors)

    def test_docx_page_count_is_bound_to_render_evidence(self) -> None:
        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        item = next(item for item in evidence["artifacts"] if item["id"] == "resume_docx")
        item["pages"] = item["pages"][:1]
        self.save(evidence_path, evidence)
        with self.assertRaisesRegex(ContractError, "resume_docx has 1 pages; expected 2"):
            run_manifest(self.manifest)

    def test_stale_review_page_hash_blocks_release(self) -> None:
        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["page_image_hashes"]["resume_docx"][0] = "0" * 64
        self.save(review_path, review)
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2, report)
        self.assert_gate_failed(report, "exact_artifact_review")

    def test_changed_rendered_page_is_detected_before_release(self) -> None:
        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        item = next(item for item in evidence["artifacts"] if item["id"] == "resume_docx")
        page_path = evidence_path.parent / item["pages"][0]["path"]
        page_path.write_bytes(page_path.read_bytes() + b"changed")
        with self.assertRaisesRegex(ContractError, "Rendered page hash changed"):
            run_manifest(self.manifest)

    def test_frozen_page_mutation_is_detected_by_validate(self) -> None:
        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 0, report)
        page_path = self.run_dir / "audit" / "render-evidence" / "resume_docx" / "page-1.png"
        page_path.write_bytes(page_path.read_bytes() + b"changed")
        report = validate_run(self.run_dir)
        self.assert_gate_failed(report, "layout_quality")

    def test_unrelated_blank_page_is_rejected_even_when_all_supplied_hashes_match(self) -> None:
        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        item = next(item for item in evidence["artifacts"] if item["id"] == "resume_pdf")
        page = item["pages"][0]
        page_path = evidence_path.parent / page["path"]
        Image.new("RGB", (120, 160), (0, 0, 0)).save(page_path)
        page["sha256"] = sha256_file(page_path)
        page["bytes"] = page_path.stat().st_size
        self.save(evidence_path, evidence)

        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["page_image_hashes"]["resume_pdf"][0] = page["sha256"]
        self.save(review_path, review)

        with self.assertRaisesRegex(ContractError, "does not match a fresh render"):
            run_manifest(self.manifest)

    def test_changed_renderer_identity_blocks_release(self) -> None:
        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        evidence["tool_capabilities"]["pdftoppm"]["sha256"] = "0" * 64
        self.save(evidence_path, evidence)
        with self.assertRaisesRegex(ContractError, "renderer identity changed"):
            run_manifest(self.manifest)

    def test_default_python_docx_package_metadata_blocks_release(self) -> None:
        docx_path = self.root / "artifacts" / "resume.docx"
        document = Document()
        for line in ["Avery Candidate", "Platform protected achievement", "Independent research"]:
            document.add_paragraph(line)
        document.save(docx_path)

        evidence_path = self.root / "stages" / "render-evidence.json"
        evidence = self.load(evidence_path)
        item = next(item for item in evidence["artifacts"] if item["id"] == "resume_docx")
        item["sha256"] = sha256_file(docx_path)
        item["bytes"] = docx_path.stat().st_size
        rendered = render_exact_artifact(
            docx_path,
            "docx",
            evidence["tool_capabilities"],
            self.root / "stages" / "render-pages" / "resume_docx",
        )
        item["pages"] = [
            {
                "page": page_number,
                "path": str(path.relative_to(self.root / "stages")),
                "sha256": sha256_file(path),
                "bytes": path.stat().st_size,
            }
            for page_number, path in enumerate(rendered, start=1)
        ]
        self.save(evidence_path, evidence)

        review_path = self.root / "stages" / "review-record.json"
        review = self.load(review_path)
        review["artifact_hashes"]["resume_docx"] = item["sha256"]
        review["page_image_hashes"]["resume_docx"] = [page["sha256"] for page in item["pages"]]
        self.save(review_path, review)

        code, report = run_manifest(self.manifest)
        self.assertEqual(code, 2, report)
        self.assert_gate_failed(report, "layout_quality")
        errors = next(gate["errors"] for gate in report["gates"] if gate["id"] == "layout_quality")
        self.assertTrue(any("python-docx" in error or "thumbnail" in error for error in errors), errors)


if __name__ == "__main__":
    unittest.main()
